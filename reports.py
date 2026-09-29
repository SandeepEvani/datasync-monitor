"""Excel report generation for DataSync monitoring dashboard."""

from __future__ import annotations

from io import BytesIO
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
from openpyxl.utils import get_column_letter

from db import execution_summaries, file_detail


_BYTES_TO_MB = 1 / (1024 * 1024)
_BYTES_TO_GB = 1 / (1024 ** 3)
_BYTES_TO_TB = 1 / (1024 ** 4)

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_BYTE_HEADER_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

_SUCCESS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
_SUCCESS_FONT = Font(name="Calibri", color="375623")
_ERROR_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
_ERROR_FONT = Font(name="Calibri", color="C62828")

_TRANSFERRED_FILL = PatternFill(start_color="DBEEF4", end_color="DBEEF4", fill_type="solid")
_TRANSFERRED_FONT = Font(name="Calibri", color="1F4E79")
_VERIFIED_FILL = PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid")
_VERIFIED_FONT = Font(name="Calibri", color="6C3483")
_SKIPPED_FILL = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
_SKIPPED_FONT = Font(name="Calibri", color="F57F17")
_DELETED_FILL = PatternFill(start_color="F3E5F5", end_color="F3E5F5", fill_type="solid")
_DELETED_FONT = Font(name="Calibri", color="7B1FA2")

_BYTE_COL_FILL = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")
_ALT_ROW_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

_THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
)
_HEADER_BORDER = Border(
    bottom=Side(style="medium", color="1F4E79"),
)

_DATA_FONT = Font(name="Calibri", size=10)
_DATA_ALIGNMENT = Alignment(vertical="center", wrap_text=False)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUMBER_FORMAT = "#,##0"
_DECIMAL_FORMAT = "#,##0.000"


def _add_byte_conversions(df: pd.DataFrame, raw_col: str, label: str) -> None:
    idx = df.columns.get_loc(raw_col) + 1
    for suffix, factor in [("MB", _BYTES_TO_MB), ("GB", _BYTES_TO_GB), ("TB", _BYTES_TO_TB)]:
        col_name = f"{label} ({suffix})"
        df.insert(idx, col_name, pd.to_numeric(df[raw_col], errors="coerce").fillna(0) * factor)
        df[col_name] = df[col_name].round(3)
        idx += 1


def _is_byte_col(col_name: str) -> bool:
    return col_name.endswith(("(MB)", "(GB)", "(TB)"))


def _format_worksheet(ws, df: pd.DataFrame, sheet_type: str = "summary") -> None:
    if ws.max_row < 2:
        return

    col_names = list(df.columns)
    byte_col_indices = {i + 1 for i, c in enumerate(col_names) if _is_byte_col(c)}
    status_col = None
    report_type_col = None
    error_code_col = None

    for i, c in enumerate(col_names):
        if c == "Status":
            status_col = i + 1
        elif c == "Report Type":
            report_type_col = i + 1
        elif c == "Error Code":
            error_code_col = i + 1

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = _HEADER_FONT
        cell.fill = _BYTE_HEADER_FILL if col_idx in byte_col_indices else _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _HEADER_BORDER

    for row_idx in range(2, ws.max_row + 1):
        is_alt = row_idx % 2 == 0
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = _DATA_FONT
            cell.alignment = _DATA_ALIGNMENT
            cell.border = _THIN_BORDER

            if col_idx in byte_col_indices:
                cell.fill = _BYTE_COL_FILL
                cell.number_format = _DECIMAL_FORMAT
            elif is_alt:
                cell.fill = _ALT_ROW_FILL

            if isinstance(cell.value, (int, float)) and col_idx not in byte_col_indices:
                cell.number_format = _NUMBER_FORMAT

        if status_col:
            cell = ws.cell(row=row_idx, column=status_col)
            val = str(cell.value or "").upper()
            if val in ("COMPLETED", "SUCCESS"):
                cell.fill = _SUCCESS_FILL
                cell.font = Font(name="Calibri", bold=True, color="375623", size=10)
            elif val in ("ERROR", "FAILED"):
                cell.fill = _ERROR_FILL
                cell.font = Font(name="Calibri", bold=True, color="C62828", size=10)

        if report_type_col:
            cell = ws.cell(row=row_idx, column=report_type_col)
            val = str(cell.value or "").lower()
            if val == "transferred":
                cell.fill, cell.font = _TRANSFERRED_FILL, _TRANSFERRED_FONT
            elif val == "verified":
                cell.fill, cell.font = _VERIFIED_FILL, _VERIFIED_FONT
            elif val == "skipped":
                cell.fill, cell.font = _SKIPPED_FILL, _SKIPPED_FONT
            elif val == "deleted":
                cell.fill, cell.font = _DELETED_FILL, _DELETED_FONT

        if error_code_col:
            cell = ws.cell(row=row_idx, column=error_code_col)
            if cell.value:
                cell.fill = _ERROR_FILL
                cell.font = _ERROR_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
        max_len = 0
        for cell in col_cells:
            val = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    ws.sheet_properties.tabColor = "1F4E79" if sheet_type == "summary" else "2E75B6"


def _build_summary_df(conn, task_name: str | None = None) -> pd.DataFrame:
    df = execution_summaries(conn, task_name=task_name)
    if df.empty:
        return df

    cols_to_drop = [c for c in ["bytes_transferred", "bytes_compressed"] if c in df.columns]
    if cols_to_drop:
        df = df.drop(columns=cols_to_drop)

    rename = {
        "task_name": "Task Name",
        "task_id": "Task ID",
        "execution_id": "Execution ID",
        "task_mode": "Task Mode",
        "overall_status": "Status",
        "start_time": "Start Time",
        "end_time": "End Time",
        "total_time": "Total Time",
        "files_transferred": "Files Transferred",
        "files_verified": "Files Verified",
        "files_skipped": "Files Skipped",
        "files_deleted": "Files Deleted",
        "bytes_written": "Bytes Written",
        "prepare_duration": "Prepare Duration",
        "prepare_status": "Prepare Status",
        "transfer_duration": "Transfer Duration",
        "transfer_status": "Transfer Status",
        "verify_duration": "Verify Duration",
        "verify_status": "Verify Status",
        "error_code": "Error Code",
        "error_detail": "Error Detail",
        "files_failed_prepare": "Files Failed Prepare",
        "files_failed_transfer": "Files Failed Transfer",
        "files_failed_verify": "Files Failed Verify",
        "files_failed_delete": "Files Failed Delete",
        "files_prepared": "Files Prepared",
        "estimated_bytes_to_transfer": "Est. Bytes to Transfer",
        "estimated_files_to_transfer": "Est. Files to Transfer",
        "source_location_type": "Source Location",
        "destination_location_type": "Dest Location",
    }
    df = df.rename(columns={k: v for k, v in rename.items() if k in df.columns})

    if "Bytes Written" in df.columns:
        _add_byte_conversions(df, "Bytes Written", "Written")
    if "Est. Bytes to Transfer" in df.columns:
        _add_byte_conversions(df, "Est. Bytes to Transfer", "Est. Transfer")

    return df


def _build_detail_df(conn, task_name: str) -> pd.DataFrame:
    df = file_detail(conn, task_name=task_name, limit=1_000_000, offset=0)
    if df.empty:
        return df

    rename = {
        "task_name": "Task Name",
        "execution_id": "Execution ID",
        "report_type": "Report Type",
        "task_mode": "Task Mode",
        "relative_path": "Relative Path",
        "content_size": "Content Size",
        "item_type": "Item Type",
        "status": "Status",
        "transfer_type": "Transfer Type",
        "event_timestamp": "Event Timestamp",
        "error_code": "Error Code",
        "error_detail": "Error Detail",
    }
    df = df.rename(columns={k: v for k, v in rename.items() if k in df.columns})

    if "Content Size" in df.columns:
        _add_byte_conversions(df, "Content Size", "Size")

    return df


def _sort_df(df: pd.DataFrame, sort_cols: list[tuple[str, bool]]) -> pd.DataFrame:
    cols = [c for c, _ in sort_cols if c in df.columns]
    asc = [a for c, a in sort_cols if c in df.columns]
    if cols:
        df = df.sort_values(cols, ascending=asc).reset_index(drop=True)
    return df


def generate_all_tasks_report(conn) -> tuple[bytes, str]:
    df = _build_summary_df(conn)
    df = _sort_df(df, [("Task Name", True), ("Start Time", False)])
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        sheet_name = "All Tasks - Summary"
        if df.empty:
            pd.DataFrame({"Info": ["No summary data available"]}).to_excel(
                writer, sheet_name=sheet_name, index=False
            )
        else:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_worksheet(writer.sheets[sheet_name], df, "summary")

    filename = f"all_tasks_report_{date.today().isoformat()}.xlsx"
    return buf.getvalue(), filename


def generate_task_report(conn, task_name: str) -> tuple[bytes, str]:
    df_summary = _build_summary_df(conn, task_name=task_name)
    df_summary = _sort_df(df_summary, [("Start Time", False)])
    df_detail = _build_detail_df(conn, task_name)
    df_detail = _sort_df(df_detail, [
        ("Execution ID", True), ("Report Type", True), ("Event Timestamp", False),
    ])

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        if df_summary.empty:
            pd.DataFrame({"Info": ["No summary data available"]}).to_excel(
                writer, sheet_name="Summary", index=False
            )
        else:
            df_summary.to_excel(writer, sheet_name="Summary", index=False)
            _format_worksheet(writer.sheets["Summary"], df_summary, "summary")

        if df_detail.empty:
            pd.DataFrame({"Info": ["No detailed file records available"]}).to_excel(
                writer, sheet_name="Detailed", index=False
            )
        else:
            df_detail.to_excel(writer, sheet_name="Detailed", index=False)
            _format_worksheet(writer.sheets["Detailed"], df_detail, "detail")

    safe_name = task_name.replace("/", "_").replace("\\", "_")
    filename = f"{safe_name}_report_{date.today().isoformat()}.xlsx"
    return buf.getvalue(), filename
